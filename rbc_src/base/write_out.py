# -*- coding: utf-8 -*-
import sys, json, csv, os, zipfile
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
OUT = sys.argv[1]; DEST = sys.argv[2]
d = json.load(open(f'{OUT}/final.json'))
os.makedirs(f'{DEST}/csv', exist_ok=True)

RUBRICAS = [
 ('VB', 'Vencimento básico do cargo (ou da referência)'), ('GAJ', 'Gratificação de Atividade Judiciária'), ('APJ', 'Adicional de Padrão Judiciário (Lei 9.421/96)'),
 ('DIF_28_86', 'Diferença de 28,86% (Leis 8.622 e 8.627/93)'), ('GRAT_EXTRAORDINARIA', 'Gratificação Extraordinária (170%)'), ('GRAT_JUDICIARIA', 'Gratificação Judiciária (80%) — 1993'),
 ('ABONO', 'Abono de 5% sobre vencimento (MP 433/94)'), ('REDUTOR', 'Redutor (parcela subtraída do total bruto; valor positivo = a deduzir)'),
 ('TOTAL_BRUTO', 'Soma das parcelas antes do redutor'), ('TOTAL', 'Total da linha conforme a planilha (subtotal; não é rubrica de folha)'),
 ('GAS', 'Gratificação de Atividade de Segurança'), ('GAE', 'Gratificação de Atividade Externa'),
 ('SUBSIDIO', 'Subsídio de magistrado'), ('SUBSTITUICAO', 'Diferença de subsídio por substituição'), ('PERC_REPRESENTACAO', 'Percentual (ou fator) da representação mensal sobre o vencimento — não é valor monetário'),
 ('REPRESENTACAO_MENSAL', 'Representação mensal (magistrados, até 2004)'), ('PAE', 'Parcela Autônoma de Equivalência'), ('PARCELA_AUTONOMA_EQUIVALENCIA', 'Parcela Autônoma de Equivalência'),
 ('PAE_AUXILIO_MORADIA', 'Parcela da PAE referente a auxílio-moradia'), ('PAE_AUX_MORADIA_ESCALONADO_5PCT', 'Parcela escalonada (5%) do auxílio-moradia na PAE'),
 ('PARCELA_ATO_TST_GP_109_2001', 'Parcela do Ato TST.GP 109/2001'), ('PARCELA_STF_195_2000', 'Parcela STF 195/00'), ('VALOR_POR_SESSAO', 'Juiz classista: valor por sessão a que comparecer'),
 ('FC_OPCAO_70PCT', 'Função comissionada para quem optou pela remuneração do cargo efetivo (70%)'), ('FC_OPCAO_65PCT', 'CJ para quem optou pela remuneração do cargo efetivo (65%)'),
 ('FC_INTEGRAL', 'Função/cargo em comissão — valor integral (não optante)'), ('FC_VALOR', 'Valor da FC (FC não tem opção; valor único)'),
 ('FC_INTEGRAL_DERIVADO', 'NÃO USAR — FC ÷ 0,65 calculado na planilha sem base legal'), ('FC_VALOR_BASE', 'Valor-base da FC (Lei 9.421/96)'), ('FC_PERC_GAJ', 'Percentual da GAJ sobre a FC'),
 ('FC_GAJ', 'GAJ incidente sobre a FC'), ('FC_APJ', 'APJ incidente sobre a FC'), ('FC_TOTAL', 'Total da FC integral'), ('FC_TOTAL_BRUTO', 'Total da FC antes do redutor'),
 ('DECIMOS', 'Décimos incorporados (fração n/10 da FC)'), ('DECIMOS_VALOR', 'Décimos — valor antes do redutor'), ('DECIMOS_REDUTOR', 'Décimos — redutor'), ('DECIMOS_TOTAL', 'Décimos — valor após redutor'),
 ('VPNI', 'Vantagem Pessoal Nominalmente Identificada (quintos/décimos) — fração'), ('VB_REF_VPI', 'Vencimento de referência para a VPI 13,23%'), ('VPI_12_23', 'VPI (rótulo da planilha: 12,23%)'),
 ('GAJ_SOBRE_VPI', 'GAJ sobre a VPI'), ('FC_CC_REF_VPI', 'FC/CC de referência para a VPI'),
 ('DAS_VENCIMENTO', 'DAS — vencimento (ou retribuição única para DAS-04 a 06 a partir da Lei 9.030/95)'), ('DAS_DIFERENCA', 'DAS — diferença'), ('DAS_SUBTOTAL', 'DAS — subtotal'),
 ('DAS_PERC_REPRESENTACAO', 'DAS — fator de representação'), ('DAS_REPRESENTACAO_MENSAL', 'DAS — representação mensal'), ('DAS_55PCT', 'DAS — 55%'), ('GADF_PERC', 'GADF — fator'),
 ('GADF', 'Gratificação de Atividade pelo Desempenho de Função'), ('GADF_55PCT', 'GADF — 55%'), ('GRAT_GABINETE', 'Gratificação de gabinete'), ('DIFERENCA', 'Diferença'), ('SOMA', 'Soma parcial')]
DICT = [
 ('valores', 'id', 'Identificador único do registro (V00001…)'), ('valores', 'grupo', 'SERVIDOR_EFETIVO, MAGISTRADO, FUNCAO_COMISSIONADA, DECIMOS_QUINTOS, DAS, GRATIFICACAO_GABINETE, VPI_13_23, VPNI_QUINTOS'),
 ('valores', 'vigencia_inicio / vigencia_fim', 'Período em que o valor vale (fim vazio = sem fim informado)'), ('valores', 'moeda', 'NCz$, Cr$, CR$, URV, R$ — valor NOMINAL, sem conversão'),
 ('valores', 'cargo', 'Cargo como consta na planilha (até 1994: NÍVEL AUXILIAR/INTERMEDIÁRIO/SUPERIOR)'), ('valores', 'ref / classe / padrao', 'Referência original e sua decomposição (classe A–D; padrão numérico ou romano)'),
 ('valores', 'funcao / modalidade', 'FC/CJ/DAS e modalidade: OPTANTE, INTEGRAL, UNICO, COM/SEM VÍNCULO'), ('valores', 'fracao', 'Fração de décimos/quintos (ex.: 3/10, 2/5)'),
 ('valores', 'rubrica', 'Código da parcela — ver aba rubricas'), ('valores', 'unidade', 'R$/moeda da época, % ou fator'), ('valores', 'e_total', 'VERDADEIRO = subtotal da planilha, não somar com as parcelas'),
 ('valores', 'status', 'ok | verificar (falhou validação) | nao_usar (valor sem base legal)'), ('valores', 'aba_origem / celula_origem', 'Rastreabilidade até a planilha original'),
 ('valores', 'origem_formula', 'Fórmula original, quando a célula era calculada'),
 ('competencias_servidor', '(todas)', 'Uma linha por competência (mês) × cargo × classe × padrão, de 01/1990 a 12/2026, com a tabela vigente no dia 1º. situacao: OK | MUDANCA_NO_MES (nova tabela começa no meio do mês: ver nova_tabela_em) | VERIFICAR_VALOR'),
 ('competencias_magistrado', '(todas)', 'Idem para magistrados (cargo padronizado)'), ('competencias_funcao', '(todas)', 'Idem para FC/CJ; SEM_DADO_NA_PLANILHA nos meses sem tabela'),
 ('parametros_mensais', '(todas)', 'Teto do RGPS, salário mínimo, fator da GAJ, fator de implantação GAE/GAS, VPI, percentuais do AQ (até 2025), VR do AQ (2026) e contribuição limitada ao teto, por competência')]

def ws_write(wb, name, cols, data, widths=None, numfmt=None):
    ws = wb.create_sheet(name)
    ws.append(cols)
    for c in ws[1]:
        c.font = Font(name='Arial', bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='1F3864'); c.alignment = Alignment(wrap_text=True, vertical='top')
    for row in data: ws.append([row.get(k) if isinstance(row, dict) else row[i] for i, k in enumerate(cols)])
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = Font(name='Arial', size=10)
            if numfmt and cols[c.column - 1] in numfmt and isinstance(c.value, (int, float)): c.number_format = numfmt[cols[c.column - 1]]
    ws.freeze_panes = 'A2'; ws.auto_filter.ref = ws.dimensions
    for i, k in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(k, min(max(len(k) + 2, 10), 40))
    with open(f'{DEST}/csv/{name}.csv', 'w', newline='', encoding='utf-8-sig') as f:
        w = csv.writer(f, delimiter=';'); w.writerow(cols)
        for row in data: w.writerow(['' if (row.get(k) if isinstance(row, dict) else row[i]) is None else (row.get(k) if isinstance(row, dict) else row[i]) for i, k in enumerate(cols)])
    return ws

wb = Workbook(); lm = wb.active; lm.title = 'LEIA-ME'
st = d['stats']
txt = [
 ('BASE DE REMUNERAÇÕES TRT-17 PARA CÁLCULO DE RBC', True),
 ('Origem: TABELA_REMUNERAÇÃO.xlsx (DIPROF). Gerado em 09/10/2026, com as decisões da DIPROF de 09/10/2026 sobre os alertas (coluna situacao da aba alertas). Valores NOMINAIS, como estão na planilha original. A RBC cobre todo o período de contribuição, inclusive antes de 07/1994, na moeda da época e sem conversão; a média da aposentadoria usa só de 07/1994 em diante.', False),
 ('', False),
 ('Abas', True),
 ('valores — base longa: 1 linha = 1 valor (vigência × cargo/função × referência × rubrica). Fonte de verdade.', False),
 ('competencias_servidor / _magistrado / _funcao — a mesma base expandida por mês (01/1990–12/2026; antes de 07/1994 na moeda da época), pronta para consulta da IA.', False),
 ('parametros / parametros_mensais — teto RGPS, salário mínimo, GAJ, GAE/GAS, VPI, AQ (percentuais até 2025 e VR a partir de 2026), IR (deduções).', False),
 ('tabela_ir — faixas do IRRF (2007 em diante, conforme planilha).', False),
 ('regras_incidencia — anotações da DIPROF sobre incidência de PSS (aba ANOTAÇÕES).', False),
 ('alertas — inconsistências da planilha original, com a situação após a revisão da DIPROF (resolvido, aceito, tratado, a_confirmar ou aberto). Os a_confirmar trazem uma proposta já aplicada ou uma pergunta que depende da DIPROF.', False),
 ('rubricas / dicionario — significado dos códigos e colunas.', False),
 ('', False),
 ('Validação automática', True),
 (f"{st['nchk']} totais conferidos contra a soma das parcelas: {st['nerr']} divergência(s). {st['nmono']} salto(s)/inversão(ões) de vencimento detectado(s). Todas as células numéricas não importadas foram conferidas: são duplicatas de valores já importados, cabeçalhos ou itens listados em alertas.", False),
 ('', False),
 ('Regras de uso (resumo — detalhes no guia_ia.md)', True),
 ('1. Use a tabela vigente no dia 1º da competência; se situacao = MUDANCA_NO_MES, calcule pro rata pelos dias (a partir de dez/2017 o divisor é o nº de dias do mês — Res. CSJT 211/2017; antes, 30).', False),
 ('2. Ignore registros com status = nao_usar; trate status = verificar como pendência de conferência humana.', False),
 ('3. Colunas TOTAL/SOMA/SUBTOTAL (e_total = VERDADEIRO) não devem ser somadas às parcelas.', False),
 ('4. Para meses SEM_DADO_NA_PLANILHA (ex.: FC/CJ de 2009 a 2022) a IA deve sinalizar a lacuna, não estimar.', False),
]
for i, (t, b) in enumerate(txt, 1):
    c = lm.cell(i, 1, t); c.font = Font(name='Arial', bold=b, size=12 if i == 1 else 10); c.alignment = Alignment(wrap_text=True)
lm.column_dimensions['A'].width = 140

money = {'valor': '#,##0.00####'}
ws_write(wb, 'valores', d['COLS'], d['rows'], widths={'cargo': 28, 'base_legal': 40, 'observacao': 50, 'vigencia_texto': 30, 'origem_formula': 25}, numfmt=money)
sc = ['competencia', 'cargo', 'classe', 'padrao', 'ref', 'moeda', 'vigencia_tabela'] + d['SRUB'] + ['situacao', 'nova_tabela_em', 'ids_origem']
ws_write(wb, 'competencias_servidor', sc, d['comp_serv'], widths={'cargo': 24, 'ids_origem': 40}, numfmt={k: '#,##0.00####' for k in d['SRUB']})
mc = ['competencia', 'cargo', 'moeda', 'vigencia_tabela'] + d['MRUB'] + ['situacao', 'nova_tabela_em', 'ids_origem']
ws_write(wb, 'competencias_magistrado', mc, d['comp_mag'], widths={'cargo': 38}, numfmt={k: '#,##0.00' for k in d['MRUB']})
fcc = ['competencia', 'funcao', 'modalidade', 'moeda', 'vigencia_tabela'] + d['FRUB'] + ['situacao', 'nova_tabela_em', 'ids_origem']
ws_write(wb, 'competencias_funcao', fcc, d['comp_fc'], numfmt={k: '#,##0.00####' for k in d['FRUB']})
pc = ['id', 'parametro', 'vigencia_inicio', 'vigencia_fim', 'valor', 'unidade', 'base_legal', 'observacao', 'aba_origem', 'celula_origem']
ws_write(wb, 'parametros', pc, d['params'], widths={'parametro': 34, 'observacao': 60, 'base_legal': 34})
pmc = ['competencia', 'teto_rgps', 'salario_minimo', 'gaj_fator_sobre_vb', 'gae_gas_fator_implantacao', 'vpi_lei_10698', 'aq_pct_doutorado', 'aq_pct_mestrado', 'aq_pct_especializacao', 'aq_pct_treinamento_120h', 'aq_pct_graduacao_tecnico', 'aq_valor_referencia', 'contribuicao_limitada_teto']
ws_write(wb, 'parametros_mensais', pmc, d['pm'])
irc = ['vigencia_inicio', 'vigencia_fim', 'faixa', 'base_de', 'base_ate', 'aliquota', 'parcela_deduzir', 'base_legal', 'aba_origem', 'linha_origem']
ws_write(wb, 'tabela_ir', irc, d['ir'], numfmt={'aliquota': '0.0%'})
ws_write(wb, 'regras_incidencia', ['fonte', 'texto'], d['regras'], widths={'texto': 120})
ws_write(wb, 'listas_funcoes', ['lista', 'funcao', 'nivel_fc', 'celula'], d['listas'], widths={'lista': 36, 'funcao': 40})
ws_write(wb, 'alertas', ['id', 'gravidade', 'situacao', 'tipo', 'aba', 'celula', 'descricao', 'decisao_diprof'], d['alerts'], widths={'descricao': 110, 'tipo': 30, 'decisao_diprof': 70, 'situacao': 14})
ws_write(wb, 'rubricas', ['rubrica', 'descricao'], RUBRICAS, widths={'rubrica': 34, 'descricao': 90})
ws_write(wb, 'dicionario', ['aba', 'coluna', 'descricao'], DICT, widths={'aba': 24, 'coluna': 28, 'descricao': 110})
for ws in wb.worksheets:
    if ws.title == 'alertas':
        for row in ws.iter_rows(min_row=2):
            g = row[1].value
            fill = {'alta': 'F8CBAD', 'média': 'FFE699', 'baixa': 'E2EFDA'}.get(g)
            if fill: row[1].fill = PatternFill('solid', fgColor=fill)
            row[5].alignment = Alignment(wrap_text=True, vertical='top')
wb.save(f'{DEST}/Base_Remuneracao_RBC_TRT17.xlsx')
print('ok')
